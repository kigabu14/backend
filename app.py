"""
Lazada Dashboard — Flask Backend
รัน: python app.py
"""
import json, os, io
from datetime import datetime, timedelta
from flask import Flask, jsonify, request, send_file
from flask_cors import CORS

# ── Lazada client ────────────────────────────────────────────────────────────
from lazada import LazadaClient

CONFIG_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "config.json")

def _cfg():
    with open(CONFIG_FILE, "r", encoding="utf-8") as f:
        return json.load(f)

def orders_client() -> LazadaClient:
    c = _cfg()["orders"]
    return LazadaClient(c["app_key"], c["app_secret"], c.get("access_token"))

def lazpay_client() -> LazadaClient:
    c = _cfg()["lazpay"]
    return LazadaClient(c["app_key"], c["app_secret"], c.get("access_token"))

# ── Flask app ────────────────────────────────────────────────────────────────
app = Flask(__name__)
CORS(app, resources={r"/api/*": {"origins": "*"}})

def ok(data): return jsonify({"success": True, "data": data})

@app.route("/api/health")
def health():
    return jsonify({"status": "ok", "service": "p2p-lazada-backend"})
def err(msg, code=400): return jsonify({"success": False, "error": str(msg)}), code

# ═══════════════════════════════════════════════════════════════════════════
# ORDERS
# ═══════════════════════════════════════════════════════════════════════════

@app.route("/api/orders")
def get_orders():
    try:
        status = request.args.get("status", "ready_to_ship")
        limit  = int(request.args.get("limit", 50))
        days   = int(request.args.get("days", 30))
        data   = orders_client().get_orders(status=status, limit=limit, days=days)
        orders = data.get("orders", []) if isinstance(data, dict) else []
        rows = []
        for o in orders:
            bill = o.get("address_billing") or {}
            rows.append({
                "order_id":     o.get("order_id"),
                "order_number": o.get("order_number"),
                "status":       o.get("statuses", [o.get("status", "")]),
                "price":        o.get("price"),
                "item_count":   o.get("items_count"),
                "created_at":   o.get("created_at"),
                "buyer":        f"{bill.get('first_name','')} {bill.get('last_name','')}".strip(),
            })
        return ok({"total": len(rows), "orders": rows})
    except Exception as e:
        return err(e)

@app.route("/api/orders/<int:order_id>/items")
def get_order_items(order_id):
    try:
        items = orders_client().get_order_items(order_id)
        return ok(items)
    except Exception as e:
        return err(e)

@app.route("/api/orders/summary")
def orders_summary():
    try:
        client  = orders_client()
        summary = {}
        for status in ["ready_to_ship", "pending", "shipped", "delivered", "canceled", "returned"]:
            try:
                data   = client.get_orders(status=status, limit=100, days=1)
                orders = data.get("orders", []) if isinstance(data, dict) else []
                summary[status] = len(orders)
            except:
                summary[status] = 0
        total = sum(summary.values())
        return ok({"as_of": datetime.now().strftime("%Y-%m-%d %H:%M"), "total_today": total, "by_status": summary})
    except Exception as e:
        return err(e)

# ═══════════════════════════════════════════════════════════════════════════
# FINANCE
# ═══════════════════════════════════════════════════════════════════════════

@app.route("/api/finance/transactions")
def get_transactions():
    try:
        days       = int(request.args.get("days", 7))
        page       = int(request.args.get("page", 1))
        trans_type = request.args.get("trans_type", "")
        end        = datetime.now()
        start      = end - timedelta(days=days)
        # ✅ path ถูกต้อง: /finance/transaction/details/get (มี s)
        data = lazpay_client().query_transaction_details(
            start.strftime("%Y-%m-%d 00:00:00"),
            end.strftime("%Y-%m-%d 23:59:59"),
            trans_type=trans_type,
            page=page, page_size=100
        )
        return ok(data)
    except Exception as e:
        return err(e)

@app.route("/api/finance/payout")
def get_payout():
    try:
        days  = int(request.args.get("days", 30))
        end   = datetime.now()
        start = end - timedelta(days=days)
        data  = lazpay_client().get_payout_status(
            start.strftime("%Y-%m-%d 00:00:00"),
            end.strftime("%Y-%m-%d 23:59:59")
        )
        return ok(data)
    except Exception as e:
        return err(e)

@app.route("/api/finance/summary")
def finance_summary():
    """สรุปค่าคอม/ค่าธรรมเนียม/VAT/WHT/ยอดสุทธิ"""
    try:
        days  = int(request.args.get("days", 7))
        end   = datetime.now()
        start = end - timedelta(days=days)
        data  = lazpay_client().query_transaction_details(
            start.strftime("%Y-%m-%d 00:00:00"),
            end.strftime("%Y-%m-%d 23:59:59"),
            page_size=200
        )
        txns = data if isinstance(data, list) else data.get("transactions", []) if isinstance(data, dict) else []

        summary = {
            "gross_revenue": 0, "commission": 0, "shipping_fee": 0,
            "vat": 0, "wht": 0, "returns": 0, "net_income": 0,
            "transactions": txns
        }
        for t in txns:
            amount = float(t.get("amount", 0) or 0)
            ttype  = str(t.get("transaction_type", "") or "").lower()
            if "commission"    in ttype: summary["commission"]    += abs(amount)
            elif "shipping"    in ttype: summary["shipping_fee"]  += abs(amount)
            elif "vat"         in ttype: summary["vat"]           += abs(amount)
            elif "withholding" in ttype or "wht" in ttype: summary["wht"] += abs(amount)
            elif "refund"      in ttype or "return" in ttype: summary["returns"] += abs(amount)
            elif amount > 0:             summary["gross_revenue"] += amount

        summary["net_income"] = (summary["gross_revenue"] - summary["commission"]
                                  - summary["shipping_fee"] - summary["vat"]
                                  - summary["wht"] - summary["returns"])
        return ok(summary)
    except Exception as e:
        return err(e)

# ═══════════════════════════════════════════════════════════════════════════
# PRODUCTS
# ═══════════════════════════════════════════════════════════════════════════

@app.route("/api/products")
def get_products():
    try:
        status = request.args.get("status", "active")
        limit  = int(request.args.get("limit", 50))
        search = request.args.get("search", "")
        data   = orders_client().get_products(status=status, limit=limit, search=search)
        products = data.get("products", []) if isinstance(data, dict) else []
        rows = []
        for p in products:
            skus = p.get("skus", [])
            rows.append({
                "item_id": p.get("item_id"),
                "name":    p.get("attributes", {}).get("name", ""),
                "status":  p.get("status"),
                "skus": [{
                    "sku_id":     s.get("SkuId"),
                    "seller_sku": s.get("SellerSku"),
                    "price":      s.get("price"),
                    "sale_price": s.get("special_price"),
                    "quantity":   s.get("quantity"),
                } for s in skus],
            })
        return ok({"total": len(rows), "products": rows})
    except Exception as e:
        return err(e)

@app.route("/api/products/update", methods=["POST"])
def update_product():
    try:
        body       = request.json
        seller_sku = body.get("seller_sku")
        sku        = {"SellerSku": seller_sku}
        if body.get("price")      is not None: sku["price"]         = str(body["price"])
        if body.get("quantity")   is not None: sku["quantity"]      = str(body["quantity"])
        if body.get("sale_price") is not None: sku["special_price"] = str(body["sale_price"])
        # ✅ path ถูกต้อง: /product/price_quantity/update (underscore)
        result = orders_client().update_price_quantity([sku])
        return ok(result)
    except Exception as e:
        return err(e)

# ═══════════════════════════════════════════════════════════════════════════
# REVERSE RETURNS — /order/reverse/return/detail/list
# ═══════════════════════════════════════════════════════════════════════════

@app.route("/api/returns")
def get_returns():
    try:
        days   = int(request.args.get("days", 30))
        page   = int(request.args.get("page", 1))
        end    = datetime.now()
        start  = end - timedelta(days=days)
        # ✅ ใช้ /order/reverse/return/detail/list
        data   = orders_client().get_reverse_return_list(
            start.strftime("%Y-%m-%d 00:00:00"),
            end.strftime("%Y-%m-%d 23:59:59"),
            page_size=40, page_number=page
        )
        returns = data if isinstance(data, list) else data.get("reverse_order_lines", []) if isinstance(data, dict) else []
        return ok({"total": len(returns), "returns": returns})
    except Exception as e:
        return err(e)

@app.route("/api/returns/update", methods=["POST"])
def update_return():
    """อนุมัติ/ปฏิเสธการคืนสินค้า"""
    try:
        body    = request.json
        line_id = body.get("reverse_order_line_id")
        status  = body.get("status")   # REFUND_APPROVED | REFUND_REJECTED
        remark  = body.get("remark", "")
        images  = body.get("images", [])
        if not line_id or not status:
            return err("ต้องระบุ reverse_order_line_id และ status")
        result = orders_client().update_reverse_return(line_id, status, remark, images)
        return ok(result)
    except Exception as e:
        return err(e)

# ═══════════════════════════════════════════════════════════════════════════
# FBI / SYSTEM EXPORT — /fbi/download/*
# ═══════════════════════════════════════════════════════════════════════════

@app.route("/api/fbi/export", methods=["POST"])
def fbi_export():
    """POST /fbi/download/startExportByDataset — export ข้อมูลชุดใหญ่ (async)"""
    try:
        body         = request.json or {}
        dataset_name = body.get("dataset_name", "order")
        days         = int(body.get("days", 30))
        end          = datetime.now()
        start        = end - timedelta(days=days)
        result = orders_client().fbi_start_export(
            dataset_name,
            start.strftime("%Y-%m-%d 00:00:00"),
            end.strftime("%Y-%m-%d 23:59:59"),
        )
        return ok(result)
    except Exception as e:
        return err(e)

@app.route("/api/fbi/status")
def fbi_status():
    """GET /fbi/download/getExportStatus — เช็คสถานะ export job"""
    try:
        export_id = request.args.get("export_id")
        if not export_id:
            return err("ต้องระบุ export_id")
        result = orders_client().fbi_get_export_status(export_id)
        return ok(result)
    except Exception as e:
        return err(e)

# ═══════════════════════════════════════════════════════════════════════════
# EXCEL EXPORT + GOOGLE DRIVE
# ═══════════════════════════════════════════════════════════════════════════

@app.route("/api/export/excel")
def export_excel():
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        days  = int(request.args.get("days", 7))
        end   = datetime.now()
        start = end - timedelta(days=days)

        # ดึงข้อมูล Finance
        data = lazpay_client().query_transaction_details(
            start.strftime("%Y-%m-%d 00:00:00"),
            end.strftime("%Y-%m-%d 23:59:59"),
            page_size=200
        )
        txns = data if isinstance(data, list) else data.get("transactions", []) if isinstance(data, dict) else []

        # สร้าง Excel
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Finance Report"

        # Header
        headers = ["วันที่", "Order ID", "ประเภท", "รายละเอียด", "ยอดรวม (฿)", "ค่าคอม (฿)", "ค่าขนส่ง (฿)", "VAT (฿)", "WHT (฿)", "ยอดสุทธิ (฿)"]
        header_fill = PatternFill("solid", fgColor="1E3A5F")
        header_font = Font(bold=True, color="FFFFFF")
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.fill = header_fill; cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        # Data rows
        for row_idx, t in enumerate(txns, 2):
            amount = float(t.get("amount", 0) or 0)
            ttype  = str(t.get("transaction_type", "") or "")
            ws.cell(row=row_idx, column=1, value=t.get("transaction_date", ""))
            ws.cell(row=row_idx, column=2, value=t.get("order_id", ""))
            ws.cell(row=row_idx, column=3, value=ttype)
            ws.cell(row=row_idx, column=4, value=t.get("details", t.get("feature_type", "")))
            commission = shipping = vat = wht = 0
            tl = ttype.lower()
            if   "commission"    in tl: commission = abs(amount)
            elif "shipping"      in tl: shipping   = abs(amount)
            elif "vat"           in tl: vat        = abs(amount)
            elif "withholding"   in tl or "wht" in tl: wht = abs(amount)
            gross = amount if amount > 0 and not any(x in tl for x in ["commission","shipping","vat","wht","refund"]) else 0
            net   = gross - commission - shipping - vat - wht
            for col, val in enumerate([gross, commission, shipping, vat, wht, net], 5):
                cell = ws.cell(row=row_idx, column=col, value=round(val, 2))
                if val < 0: cell.font = Font(color="CC0000")

        # Auto column width
        for col in ws.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 30)

        # Summary sheet
        ws2 = wb.create_sheet("Summary")
        ws2["A1"] = "รายงานสรุป"
        ws2["A1"].font = Font(bold=True, size=14)
        ws2["A3"] = "ช่วงเวลา"
        ws2["B3"] = f"{start.strftime('%Y-%m-%d')} ถึง {end.strftime('%Y-%m-%d')}"
        ws2["A4"] = "จำนวน Transaction"
        ws2["B4"] = len(txns)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        filename = f"lazada_finance_{start.strftime('%Y%m%d')}_{end.strftime('%Y%m%d')}.xlsx"
        return send_file(buf, as_attachment=True, download_name=filename,
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    except ImportError:
        return err("กรุณาติดตั้ง openpyxl: pip install openpyxl")
    except Exception as e:
        return err(e)

@app.route("/api/export/drive", methods=["POST"])
def export_to_drive():
    """Upload Excel ไป Google Drive"""
    try:
        from googleapiclient.discovery import build
        from googleapiclient.http import MediaIoBaseUpload
        from google.oauth2.credentials import Credentials
        import openpyxl, io

        body   = request.json or {}
        days   = int(body.get("days", 7))
        creds_data = body.get("credentials")  # Google OAuth credentials
        folder_id  = body.get("folder_id", "root")

        if not creds_data:
            return err("ต้องส่ง Google credentials มาด้วย")

        # Build Excel (เหมือน export_excel)
        end   = datetime.now()
        start = end - timedelta(days=days)
        data  = lazpay_client().query_transaction_details(
            start.strftime("%Y-%m-%d 00:00:00"),
            end.strftime("%Y-%m-%d 23:59:59"), page_size=200)
        txns = data if isinstance(data, list) else data.get("transactions", []) if isinstance(data, dict) else []

        wb = openpyxl.Workbook()
        ws = wb.active; ws.title = "Finance"
        headers = ["วันที่","Order ID","ประเภท","ยอดรวม","ค่าคอม","ค่าขนส่ง","VAT","WHT","ยอดสุทธิ"]
        for col, h in enumerate(headers, 1): ws.cell(row=1, column=col, value=h)
        for i, t in enumerate(txns, 2):
            amount = float(t.get("amount", 0) or 0)
            ws.cell(row=i, column=1, value=t.get("transaction_date",""))
            ws.cell(row=i, column=2, value=t.get("order_id",""))
            ws.cell(row=i, column=3, value=t.get("transaction_type",""))
            ws.cell(row=i, column=4, value=amount)
            ws.cell(row=i, column=9, value=amount)

        buf = io.BytesIO(); wb.save(buf); buf.seek(0)
        filename = f"lazada_finance_{start.strftime('%Y%m%d')}_{end.strftime('%Y%m%d')}.xlsx"

        # Upload to Drive
        creds   = Credentials(**creds_data)
        service = build("drive", "v3", credentials=creds)
        meta    = {"name": filename, "parents": [folder_id]}
        media   = MediaIoBaseUpload(buf, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        file    = service.files().create(body=meta, media_body=media, fields="id,webViewLink").execute()
        return ok({"file_id": file.get("id"), "url": file.get("webViewLink"), "filename": filename})
    except Exception as e:
        return err(e)

# ═══════════════════════════════════════════════════════════════════════════

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))  # Railway inject PORT อัตโนมัติ
    print(f"🚀 Lazada Dashboard Backend running on port {port}")
    app.run(host="0.0.0.0", port=port, debug=False)
